import docx
from docx.shared import Pt, Inches, RGBColor

def md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()
    
    for s in doc.styles:
        if hasattr(s, 'font'):
            s.font.name = 'Calibri'
    
    in_code = False
    for line in lines:
        raw = line.rstrip('\r\n')
        if raw.startswith('```'):
            in_code = not in_code
            continue
        
        if not raw.strip():
            continue
            
        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            r = p.add_run(raw)
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(40, 40, 40)
            continue

        if raw.startswith('# '):
            p = doc.add_paragraph()
            run = p.add_run(raw[2:].strip())
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 51, 102)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
        elif raw.startswith('## '):
            p = doc.add_paragraph()
            run = p.add_run(raw[3:].strip())
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 102, 153)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
        elif raw.startswith('### '):
            p = doc.add_paragraph()
            run = p.add_run(raw[4:].strip())
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(51, 51, 51)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
        elif raw.startswith('#### '):
            p = doc.add_paragraph()
            run = p.add_run(raw[5:].strip())
            run.font.size = Pt(11)
            run.font.bold = True
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
        elif raw.startswith('> '):
            p = doc.add_paragraph()
            run = p.add_run(raw[2:].strip())
            run.font.italic = True
            run.font.color.rgb = RGBColor(100, 100, 100)
            p.paragraph_format.left_indent = Inches(0.4)
        elif raw.startswith('- ') or raw.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            text = raw[2:].strip()
            parts = text.split('**')
            for i, part in enumerate(parts):
                r = p.add_run(part)
                if i % 2 == 1:
                    r.font.bold = True
        elif '|' in raw and not raw.startswith(' '):
            cols = [c.strip() for c in raw.strip('|').split('|')]
            if cols and not all(c.startswith('-') for c in cols):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.2)
                r = p.add_run(' | '.join(cols))
                r.font.size = Pt(9.5)
                r.font.name = 'Consolas'
        else:
            p = doc.add_paragraph()
            parts = raw.split('**')
            for i, part in enumerate(parts):
                r = p.add_run(part)
                if i % 2 == 1:
                    r.font.bold = True
            p.paragraph_format.space_after = Pt(3)

    doc.save(docx_path)
    print(f'Successfully updated: {docx_path}')

if __name__ == '__main__':
    md_to_docx('docs/GUIDE_DEPLOIEMENT_WINDOWS.md', 'docs/GUIDE_DEPLOIEMENT_WINDOWS.docx')
    md_to_docx('docs/GUIDE_DEPLOIEMENT_WINDOWS.md', 'GUIDE_DEPLOIEMENT_WINDOWS.docx')
