"""Vue consolidée multi-sites (licence RESEAU) : totaux pondérés, cloisonnement par groupe, export, pilotage à distance."""
import datetime
import io
from decimal import Decimal

import pytest
from django.urls import reverse
from model_bakery import baker
from openpyxl import load_workbook

from core.pilotage import tableau_reseau

AUJOURDHUI = datetime.date(2026, 10, 21)


def _site(groupe, nom, tarif, paiements, capacite=40):
    """Un établissement du groupe : une classe tarifée, autant d'élèves que de ``paiements`` (montant versé par élève)."""
    etab = baker.make('etablissements.Etablissement', nom=nom, groupe=groupe)
    annee = baker.make('parametres.AnneeScolaire', etablissement=etab, libelle='2026-2027', est_courante=True,
                       date_debut=datetime.date(2026, 10, 1), date_fin=datetime.date(2027, 7, 15))
    statut = baker.make('parametres.StatutEleve', etablissement=etab)
    rubrique = baker.make('parametres.RubriquePaiement', etablissement=etab)
    cycle = baker.make('parametres.Cycle', etablissement=etab, ordre=1, actif=True)
    classe = baker.make('parametres.Classe', etablissement=etab, cycle=cycle, niveau='6e', nom='6e A', capacite_max=capacite)
    baker.make('parametres.TarifScolarite', etablissement=etab, annee_scolaire=annee, cycle=cycle, classe=classe,
               statut_eleve=statut, rubrique=rubrique, montant=Decimal(tarif), actif=True)
    for i, montant in enumerate(paiements):
        ins = baker.make('inscriptions.Inscription', classe=classe, annee_scolaire=annee, statut='AFFECTE', statut_eleve=statut,
                         eleve=baker.make('inscriptions.Eleve', genre='M' if i % 2 == 0 else 'F', date_naissance=datetime.date(2014, 1, 1)))
        if montant:
            baker.make('finances.Paiement', inscription=ins, montant=Decimal(montant), rubrique=rubrique)
    return etab


@pytest.fixture
def reseau():
    groupe = baker.make('etablissements.GroupeEtablissements', nom='Groupe Wend-Kuuni')
    # Site A : 4 élèves × 50 000 = 200 000 attendus, 150 000 encaissés (75 %) ; site B : 1 élève × 100 000, rien payé (0 %)
    a = _site(groupe, 'Site A', 50000, (50000, 50000, 50000, 0))
    b = _site(groupe, 'Site B', 100000, (0,), capacite=10)
    sans_annee = baker.make('etablissements.Etablissement', nom='Site C', groupe=groupe)
    hors_groupe = _site(None, 'Autre école', 10000, (10000,))
    return {'groupe': groupe, 'a': a, 'b': b, 'sans_annee': sans_annee, 'hors_groupe': hors_groupe}


def _directeur_reseau(client, reseau, groupe=True):
    user = baker.make('accounts.User', role='DIRECTEUR_RESEAU', etablissement=reseau['a'],
                      groupe=reseau['groupe'] if groupe else None, must_change_password=False)
    client.force_login(user)
    return user


@pytest.mark.django_db
def test_totaux_ponderes(reseau):
    consolide = tableau_reseau([reseau['a'], reseau['b'], reseau['sans_annee']], aujourdhui=AUJOURDHUI)
    lignes = {l['etablissement'].nom: l for l in consolide['lignes']}
    assert lignes['Site A']['recouvrement']['taux'] == 75.0
    assert lignes['Site B']['recouvrement']['taux'] == 0.0
    assert lignes['Site C']['annee'] is None and 'recouvrement' not in lignes['Site C']
    totaux = consolide['totaux']
    assert totaux['sites'] == 2
    assert (totaux['effectifs']['total'], totaux['effectifs']['garcons'], totaux['effectifs']['filles']) == (5, 3, 2)
    assert totaux['effectifs']['taux_remplissage'] == 10.0  # 5 / (40 + 10)
    assert (totaux['recouvrement']['attendu'], totaux['recouvrement']['encaisse'], totaux['recouvrement']['reste']) == (300000, 150000, 150000)
    assert totaux['recouvrement']['taux'] == 50.0  # pondéré : 150 000 / 300 000, pas (75 + 0) / 2
    assert totaux['recouvrement']['redevables'] == 2


@pytest.mark.django_db
def test_directeur_reseau_voit_son_groupe_seulement(client, reseau):
    _directeur_reseau(client, reseau)
    reponse = client.get(reverse('etablissements:reseau_dashboard'))
    assert reponse.status_code == 200
    noms = [l['etablissement'].nom for l in reponse.context['lignes']]
    assert noms == ['Site A', 'Site B', 'Site C']
    assert reponse.context['totaux']['recouvrement']['taux'] == 50.0
    assert reponse.context['groupes'] == []  # pas de sélecteur de groupe hors Super Admin
    html = reponse.content.decode()
    assert 'Autre école' not in html
    assert f"{reverse('core:pilotage')}?etab={reseau['b'].pk}" in html


@pytest.mark.django_db
def test_directeur_reseau_sans_groupe(client, reseau):
    _directeur_reseau(client, reseau, groupe=False)
    reponse = client.get(reverse('etablissements:reseau_dashboard'))
    assert reponse.status_code == 200
    assert reponse.context['lignes'] == []
    assert 'rattaché à aucun groupe' in reponse.content.decode()


@pytest.mark.django_db
def test_super_admin_filtre_par_groupe(client, reseau):
    client.force_login(baker.make('accounts.User', role='SUPER_ADMIN', is_superuser=True, must_change_password=False))
    url = reverse('etablissements:reseau_dashboard')
    assert len(client.get(url).context['lignes']) == 4  # tous les établissements
    reponse = client.get(url + f"?groupe={reseau['groupe'].pk}")
    assert [l['etablissement'].nom for l in reponse.context['lignes']] == ['Site A', 'Site B', 'Site C']
    assert client.get(url + '?groupe=pas-un-uuid').status_code == 200


@pytest.mark.django_db
def test_export_excel(client, reseau):
    _directeur_reseau(client, reseau)
    reponse = client.get(reverse('etablissements:reseau_dashboard') + '?format=xlsx')
    assert reponse['Content-Type'].startswith('application/vnd.openxmlformats')
    ws = load_workbook(io.BytesIO(reponse.content)).active
    lignes = [[c.value for c in r] for r in ws.iter_rows()]
    assert any(l[0] == 'Site A' and l[2] == 4 and l[7:11] == [200000, 150000, 50000, 75.0] for l in lignes)
    assert any(l[0] == 'Site C' and l[1] == 'Aucune année courante' for l in lignes)
    assert any(l[0] == 'Total réseau' and l[1] == '2 site(s)' and l[7:11] == [300000, 150000, 150000, 50.0] for l in lignes)


@pytest.mark.django_db
def test_pilotage_a_distance_limite_au_groupe(client, reseau):
    _directeur_reseau(client, reseau)
    url = reverse('core:pilotage')
    reponse = client.get(url + f"?etab={reseau['b'].pk}")
    assert reponse.status_code == 200
    assert reponse.context['etablissement_pilote'] == reseau['b']
    assert reponse.context['recouvrement']['attendu'] == 100000
    assert 'Site B' in reponse.content.decode()
    # Hors groupe : retombe sur son propre établissement
    reponse = client.get(url + f"?etab={reseau['hors_groupe'].pk}")
    assert reponse.context['etablissement_pilote'] == reseau['a']
    assert reponse.context['recouvrement']['attendu'] == 200000
    assert client.get(url + '?etab=zzz').context['etablissement_pilote'] == reseau['a']


@pytest.mark.django_db
def test_directeur_simple_ne_pilote_pas_un_autre_site(client, reseau):
    client.force_login(baker.make('accounts.User', role='DIRECTEUR', etablissement=reseau['a'], must_change_password=False))
    reponse = client.get(reverse('core:pilotage') + f"?etab={reseau['b'].pk}")
    assert reponse.context['etablissement_pilote'] == reseau['a']
    assert client.get(reverse('etablissements:reseau_dashboard')).status_code == 302


@pytest.mark.django_db
def test_menu_reseau_reserve(client, reseau):
    _directeur_reseau(client, reseau)
    assert reverse('etablissements:reseau_dashboard') in client.get(reverse('core:home')).content.decode()
    client.force_login(baker.make('accounts.User', role='DIRECTEUR', etablissement=reseau['a'], must_change_password=False))
    assert reverse('etablissements:reseau_dashboard') not in client.get(reverse('core:home')).content.decode()
