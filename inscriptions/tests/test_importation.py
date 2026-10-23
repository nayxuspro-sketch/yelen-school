"""Import Excel / CSV des élèves : moteur de lecture, reconnaissance des colonnes et assistant en quatre étapes."""
import html
import io
from datetime import date, datetime

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from model_bakery import baker

from core import importation as moteur
from inscriptions import importation
from inscriptions.models import Eleve, Inscription
from inscriptions.views_import import cle_classe

pytestmark = pytest.mark.django_db


# ── Moteur générique (sans base) ─────────────────────────────────────────

def test_conversions_des_cellules():
    assert moteur.en_date(datetime(2012, 5, 12, 0, 0)) == date(2012, 5, 12)
    assert moteur.en_date('03/09/2011') == date(2011, 9, 3)
    assert moteur.en_date('2011-09-03') == date(2011, 9, 3)
    assert moteur.en_date(41041) == date(2012, 5, 12)          # numéro de série Excel
    assert moteur.en_date('') is None
    with pytest.raises(ValueError):
        moteur.en_date('12 mai')
    assert moteur.en_genre('Masculin') == 'M' and moteur.en_genre('fille') == 'F' and moteur.en_genre(None) is None
    assert moteur.en_booleen('Oui') is True and moteur.en_booleen('') is False and moteur.en_booleen('R') is True
    assert moteur.en_telephone(70123456.0) == '70123456'
    assert moteur.en_telephone('+226 76 12 34 56') == '+22676123456'
    assert moteur.separer_nom_prenom('OUEDRAOGO Abdoul Kader') == ('OUEDRAOGO', 'Abdoul Kader')
    assert moteur.separer_nom_prenom('SAWADOGO ZONGO Awa') == ('SAWADOGO ZONGO', 'Awa')
    assert moteur.separer_nom_prenom('Kabore Issa') == ('Kabore', 'Issa')
    assert moteur.en_note('12,5', 20) == (moteur.Decimal('12.5'), False)
    assert moteur.en_note('ABS', 20) == (moteur.Decimal('0'), True)
    assert moteur.en_note(None, 20) is None
    with pytest.raises(ValueError):
        moteur.en_note(25, 20)
    assert moteur.lettre_colonne(0) == 'A' and moteur.lettre_colonne(27) == 'AB'


def test_detection_des_colonnes_tolere_les_libelles_de_l_ecole():
    entetes = ['N°', 'NOM ET PRENOMS', 'Sexe', 'Date de naissance de l\'élève', 'Lieu de naiss.', 'Classe',
               'Contact parent', 'Redoublant (Oui/Non)']
    corr = moteur.detecter_colonnes(entetes, importation.CHAMPS)
    assert corr['nom_complet'] == 1 and corr['genre'] == 2 and corr['date_naissance'] == 3
    assert corr['classe'] == 5 and corr['telephone_parent'] == 6 and corr['redoublant'] == 7
    assert corr['lieu_naissance'] == 4 and 'nom' not in corr and 'prenom' not in corr
    assert importation.verifier_correspondance({'genre': 2, 'date_naissance': 3}) != []
    assert importation.verifier_correspondance(corr) == []


@pytest.mark.parametrize('libelle, attendu', [
    ('6ème A', '6ea'), ('6EME-A', '6ea'), ('Sixième A', '6ea'), ('Tle D', 'tled'), ('Terminale D', 'tled'),
    ('CM2', 'cm2'), ('1ère C', '1rec'), ('Première C', '1rec'),
])
def test_cle_de_rapprochement_des_classes(libelle, attendu):
    assert cle_classe(libelle) == attendu


# ── Assistant complet ────────────────────────────────────────────────────

@pytest.fixture
def ecole(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    etab = baker.make('etablissements.Etablissement', code='ETAB001')
    annee = baker.make('parametres.AnneeScolaire', etablissement=etab, libelle='2025-2026', est_courante=True,
                       date_debut=date(2025, 10, 1), date_fin=date(2026, 6, 30))
    cycle = baker.make('parametres.Cycle', etablissement=etab, ordre=1)
    sixieme = baker.make('parametres.Classe', etablissement=etab, cycle=cycle, nom='6ème A', niveau='6ème')
    cinquieme = baker.make('parametres.Classe', etablissement=etab, cycle=cycle, nom='5ème B', niveau='5ème')
    return {'etab': etab, 'annee': annee, 'sixieme': sixieme, 'cinquieme': cinquieme}


@pytest.fixture
def directeur(client, ecole):
    client.force_login(baker.make('accounts.User', etablissement=ecole['etab'], role='DIRECTEUR',
                                  email='directeur@ecole.bf', must_change_password=False))
    return client


def _xlsx(lignes, titre='LISTE DES ÉLÈVES — 6ème A — 2025-2026'):
    """Classeur « tel que l'école le tient » : titre fusionné, ligne vide, puis le tableau."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Liste'
    ws.append([titre])
    ws.append([])
    for ligne in lignes:
        ws.append(ligne)
    tampon = io.BytesIO()
    wb.save(tampon)
    return SimpleUploadedFile('liste_6eA.xlsx', tampon.getvalue(),
                              content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


LIGNES_ECOLE = [
    ['N°', 'NOM ET PRENOMS', 'SEXE', 'DATE DE NAISSANCE', 'LIEU DE NAISSANCE', 'CLASSE', 'CONTACT PARENT', 'REDOUBLANT'],
    [1, 'OUEDRAOGO Abdoul Kader', 'M', datetime(2012, 5, 12), 'Ouagadougou', '6EME A', 70123456, 'Non'],
    [2, 'SAWADOGO Awa Marie', 'F', '03/09/2011', 'Koudougou', '6ème A', '+226 76 12 34 56', 'Oui'],
    [3, 'KABORE Issa', 'M', None, 'Ouagadougou', '6EME A', None, None],                  # date manquante
    [4, 'OUEDRAOGO Abdoul Kader', 'M', datetime(2012, 5, 12), 'Ouagadougou', '6EME A', None, None],  # doublon
    [5, 'ZONGO Fati', 'F', datetime(2012, 1, 1), 'Bobo-Dioulasso', 'CM2', None, 'Non'],  # classe inconnue
]


def _jusqu_a_la_verification(client, ecole, fichier):
    reponse = client.post(reverse('inscriptions:eleve_import'), {'fichier': fichier, 'annee': ecole['annee'].pk})
    assert reponse.status_code == 302 and reponse.url == reverse('inscriptions:eleve_import_correspondance')
    page = client.get(reverse('inscriptions:eleve_import_correspondance'))
    assert page.status_code == 200
    donnees = {'feuille': page.context['feuille'].nom, 'valider': '1'}
    for ligne in page.context['champs']:
        donnees[f'col_{ligne["champ"].code}'] = '' if ligne['index'] is None else str(ligne['index'])
    return page, donnees


def test_import_complet_depuis_le_fichier_de_l_ecole(directeur, ecole):
    page, donnees = _jusqu_a_la_verification(directeur, ecole, _xlsx(LIGNES_ECOLE))
    assert page.context['feuille'].ligne_entete == 3 and page.context['nb_detectes'] == 7
    assert directeur.post(reverse('inscriptions:eleve_import_correspondance'), donnees).status_code == 302

    verification = directeur.get(reverse('inscriptions:eleve_import_verification'))
    assert verification.status_code == 200
    analyse = verification.context['analyse']
    assert analyse.nb_nouveaux == 3 and analyse.nb_existants == 0
    motifs = {numero: motif for numero, _, motif in analyse.erreurs}
    assert motifs == {6: 'date de naissance manquante', 7: 'doublon de la ligne 4 dans le fichier'}
    classes = {l['libelle']: (l['position'], l['classe']) for l in verification.context['lignes_classes']}
    assert classes['6EME A'][1] == ecole['sixieme'] and classes['6ème A'][1] == ecole['sixieme']
    assert classes['CM2'][1] is None and verification.context['nb_libelles_sans_classe'] == 1
    contenu = html.unescape(verification.content.decode())
    assert 'Importer 3 élèves (les 2 lignes en erreur seront ignorées)' in contenu and 'déjà enregistré' not in contenu
    assert Eleve.objects.count() == 0                                   # rien d'écrit avant confirmation

    choix = {f'classe_{position}': str(classe.pk) for position, classe in classes.values() if classe}
    choix[f'classe_{classes["CM2"][0]}'] = str(ecole['cinquieme'].pk)
    reponse = directeur.post(reverse('inscriptions:eleve_import_verification'), choix)
    assert reponse.status_code == 302 and reponse.url == reverse('inscriptions:eleve_import_resultat')

    resultat = directeur.get(reverse('inscriptions:eleve_import_resultat'))
    assert resultat.context['resultat'] == {'crees': 3, 'inscrits': 3, 'existants': 0, 'deja_inscrits': 0,
                                            'sans_classe': 0, 'nb_erreurs': 2}
    assert 'Rapport Excel' in resultat.content.decode() and 'doublon de la ligne 4' in resultat.content.decode()
    awa = Eleve.objects.get(nom='SAWADOGO')
    assert awa.prenom == 'Awa Marie' and awa.genre == 'F' and awa.date_naissance == date(2011, 9, 3)
    assert awa.telephone_parent == '+22676123456' and awa.matricule.startswith('ETAB001-')
    assert awa.inscriptions.get().est_redoublant is True and awa.last_classe == '6ème A'
    assert Eleve.objects.get(nom='OUEDRAOGO').telephone_parent == '70123456'
    assert Inscription.objects.filter(classe=ecole['cinquieme'], eleve__nom='ZONGO').exists()
    rapport = directeur.get(reverse('inscriptions:eleve_import_rapport'))
    assert rapport.status_code == 200 and rapport['Content-Type'].endswith('sheet')
    lignes = list(openpyxl.load_workbook(io.BytesIO(rapport.content)).active.iter_rows(values_only=True))
    assert lignes[-1] == (7, 'OUEDRAOGO Abdoul Kader', 'doublon de la ligne 4 dans le fichier')


def test_colonne_quartier_alimente_le_lieu_de_residence(directeur, ecole):
    """Province / Commune / Village ont disparu : « Quartier » (village, secteur, résidence…) remplit le lieu de résidence."""
    codes = {c.code for c in importation.CHAMPS}
    assert 'lieu_residence' in codes and not codes & {'province', 'commune', 'village'}
    lignes = [['NOM ET PRENOMS', 'SEXE', 'DATE DE NAISSANCE', 'CLASSE', 'QUARTIER', 'ADRESSE PARENTS'],
              ['TRAORE Moussa', 'M', datetime(2012, 2, 2), '6EME A', 'Secteur 30 (Karpala)', 'Porte 125']]
    _, donnees = _jusqu_a_la_verification(directeur, ecole, _xlsx(lignes))
    assert donnees['col_lieu_residence'] == '4' and donnees['col_tuteur_adresse'] == '5'
    assert directeur.post(reverse('inscriptions:eleve_import_correspondance'), donnees).status_code == 302
    assert directeur.post(reverse('inscriptions:eleve_import_verification'), {}).status_code == 302
    eleve = Eleve.objects.get(nom='TRAORE')
    assert eleve.lieu_residence == 'Secteur 30 (Karpala)' and eleve.tuteur_adresse == 'Porte 125'


def test_reimport_du_meme_fichier_ne_duplique_rien(directeur, ecole):
    for _ in range(2):
        _, donnees = _jusqu_a_la_verification(directeur, ecole, _xlsx(LIGNES_ECOLE[:3]))
        directeur.post(reverse('inscriptions:eleve_import_correspondance'), donnees)
        directeur.post(reverse('inscriptions:eleve_import_verification'), {'classe_0': str(ecole['sixieme'].pk),
                                                                           'classe_1': str(ecole['sixieme'].pk)})
    resultat = directeur.get(reverse('inscriptions:eleve_import_resultat')).context['resultat']
    assert resultat['crees'] == 0 and resultat['existants'] == 2 and resultat['deja_inscrits'] == 2
    assert Eleve.objects.count() == 2 and Inscription.objects.count() == 2


def test_eleve_existant_reconnu_par_matricule_puis_inscrit_sans_nouvelle_fiche(directeur, ecole, settings):
    ancien = baker.make('inscriptions.Eleve', matricule='OLD-001', nom='TRAORE', prenom='Moussa', genre='M',
                        date_naissance=date(2011, 2, 2))
    autre = baker.make('inscriptions.Eleve', matricule='OLD-002', nom='DIALLO', prenom='Sara', genre='F',
                       date_naissance=date(2011, 3, 3))
    csv = ('Matricule;Nom;Prénom;Sexe;Né(e) le;Classe\n'
           'OLD-001;Traoré;Moussa;M;02/02/2011;6ème A\n'          # même matricule, casse/accents différents
           'OLD-002;KONE;Ali;M;04/04/2011;6ème A\n'                # matricule d'un autre élève → refusé
           ';Diallo;Sara;F;03/03/2011;6ème A\n').encode('cp1252')  # reconnue par identité, sans matricule
    fichier = SimpleUploadedFile('eleves.csv', csv, content_type='text/csv')
    page, donnees = _jusqu_a_la_verification(directeur, ecole, fichier)
    assert page.context['feuilles'] == ['CSV'] and donnees['col_matricule'] == '0' and donnees['col_nom'] == '1'
    directeur.post(reverse('inscriptions:eleve_import_correspondance'), donnees)
    verification = directeur.get(reverse('inscriptions:eleve_import_verification'))
    analyse = verification.context['analyse']
    assert analyse.nb_existants == 2 and analyse.nb_nouveaux == 0
    assert [e[2] for e in analyse.erreurs] == ['le matricule « OLD-002 » est déjà celui de DIALLO Sara : '
                                               'corrigez le matricule ou le nom']
    directeur.post(reverse('inscriptions:eleve_import_verification'), {'classe_0': str(ecole['sixieme'].pk)})
    assert Eleve.objects.count() == 2
    assert {i.eleve_id for i in Inscription.objects.filter(classe=ecole['sixieme'])} == {ancien.pk, autre.pk}
    assert ancien.inscriptions.get().annee_scolaire == ecole['annee']


def test_modele_vide_telecharge_est_reconnu_a_cent_pour_cent(directeur, ecole):
    modele = directeur.get(reverse('inscriptions:eleve_import_template'))
    assert modele.status_code == 200
    fichier = SimpleUploadedFile('modele_import_eleves.xlsx', modele.content)
    page, donnees = _jusqu_a_la_verification(directeur, ecole, fichier)
    attendus = {c.code for c in importation.CHAMPS} - {'nom_complet'}
    assert {code for code in attendus if donnees[f'col_{code}']} == attendus
    assert page.context['feuille'].lignes[0][1] == 'OUEDRAOGO'          # la ligne d'exemple est lue


def test_fichier_refuse_ou_sans_entete(directeur, ecole):
    reponse = directeur.post(reverse('inscriptions:eleve_import'),
                             {'fichier': SimpleUploadedFile('vieux.xls', b'xx'), 'annee': ecole['annee'].pk},
                             follow=True)
    assert 'Format non pris en charge' in reponse.content.decode()
    sans_entete = _xlsx([[1, 2, 3], [4, 5, 6]])
    reponse = directeur.post(reverse('inscriptions:eleve_import'), {'fichier': sans_entete, 'annee': ecole['annee'].pk})
    reponse = directeur.get(reponse.url, follow=True)
    assert "Aucune ligne d'entête trouvée" in html.unescape(reponse.content.decode())
    assert directeur.get(reverse('inscriptions:eleve_import_verification')).url == reverse('inscriptions:eleve_import')


def test_fichier_supprime_apres_import_et_acces_enseignant_refuse(client, directeur, ecole, tmp_path):
    _, donnees = _jusqu_a_la_verification(directeur, ecole, _xlsx(LIGNES_ECOLE[:2]))
    assert list((tmp_path / 'imports').glob('*.xlsx'))
    directeur.post(reverse('inscriptions:eleve_import_correspondance'), donnees)
    directeur.post(reverse('inscriptions:eleve_import_verification'), {'classe_0': str(ecole['sixieme'].pk)})
    assert not list((tmp_path / 'imports').glob('*.xlsx'))

    client.logout()
    client.force_login(baker.make('accounts.User', etablissement=ecole['etab'], role='ENSEIGNANT',
                                  email='prof@ecole.bf', must_change_password=False))
    assert client.get(reverse('inscriptions:eleve_import')).status_code == 302
