"""Statistiques officielles (rentrée, âges, fin d'année, examens) : calculs, écran, export Excel, droits."""
import datetime
from io import BytesIO

import openpyxl
import pytest
from django.urls import reverse
from model_bakery import baker

from documents.statistiques_officielles import rang_niveau, statistiques_fin_annee, statistiques_rentree


def test_ordre_officiel_des_niveaux():
    niveaux = ['Terminale', '3ème', 'CM2', '6ème', 'CP1', '2nde', '1ère', 'Atelier']
    assert sorted(niveaux, key=rang_niveau) == ['CP1', 'CM2', '6ème', '3ème', '2nde', '1ère', 'Terminale', 'Atelier']


@pytest.fixture
def ecole():
    """2025-2026 : 6ème (2 divisions) et 3ème ; une année antérieure pour distinguer les nouveaux entrants."""
    etab = baker.make('etablissements.Etablissement', nom='LE SAVOIR')
    precedente = baker.make('parametres.AnneeScolaire', etablissement=etab, libelle='2024-2025',
                            date_debut=datetime.date(2024, 10, 1), date_fin=datetime.date(2025, 6, 30))
    annee = baker.make('parametres.AnneeScolaire', etablissement=etab, libelle='2025-2026', est_courante=True,
                       date_debut=datetime.date(2025, 10, 1), date_fin=datetime.date(2026, 6, 30))
    cycle = baker.make('parametres.Cycle', etablissement=etab, nom='Post-primaire', ordre=1)
    sixieme_a = baker.make('parametres.Classe', etablissement=etab, cycle=cycle, niveau='6ème', nom='6ème A')
    sixieme_b = baker.make('parametres.Classe', etablissement=etab, cycle=cycle, niveau='6ème', nom='6ème B')
    troisieme = baker.make('parametres.Classe', etablissement=etab, cycle=cycle, niveau='3ème', nom='3ème',
                           est_classe_examen=True, type_examen='BEPC')

    def eleve(genre, naissance, **champs):
        return baker.make('inscriptions.Eleve', genre=genre, date_naissance=naissance, **champs)

    ancien = eleve('M', datetime.date(2013, 5, 2))
    baker.make('inscriptions.Inscription', eleve=ancien, classe=sixieme_a, annee_scolaire=precedente, statut='AFFECTE')
    inscriptions = {
        'redoublant': baker.make('inscriptions.Inscription', eleve=ancien, classe=sixieme_a, annee_scolaire=annee, statut='AFFECTE', est_redoublant=True),
        'nouvelle': baker.make('inscriptions.Inscription', eleve=eleve('F', datetime.date(2014, 1, 15)), classe=sixieme_b, annee_scolaire=annee, statut='AFFECTE'),
        'abandon': baker.make('inscriptions.Inscription', eleve=eleve('F', datetime.date(2014, 7, 1)), classe=sixieme_b, annee_scolaire=annee, statut='ABANDON'),
        'candidat': baker.make('inscriptions.Inscription', eleve=eleve('M', datetime.date(2010, 3, 9)), classe=troisieme, annee_scolaire=annee, statut='AFFECTE'),
        'candidate': baker.make('inscriptions.Inscription', eleve=eleve('F', datetime.date(2011, 11, 30)), classe=troisieme, annee_scolaire=annee, statut='AFFECTE'),
    }
    conseil = baker.make('viescolaire.ConseilClasse', classe=troisieme, date_conseil=datetime.date(2026, 6, 20))
    baker.make('viescolaire.DecisionConseil', conseil=conseil, inscription=inscriptions['candidat'], decision='PASSAGE')
    baker.make('viescolaire.DecisionConseil', conseil=conseil, inscription=inscriptions['candidate'], decision='RACHAT')
    conseil_6 = baker.make('viescolaire.ConseilClasse', classe=sixieme_a, date_conseil=datetime.date(2026, 6, 21))
    baker.make('viescolaire.DecisionConseil', conseil=conseil_6, inscription=inscriptions['redoublant'], decision='REDOUBLEMENT')
    session = baker.make('examens.SessionExamen', annee_scolaire=annee, type_examen='BEPC', libelle='BEPC 2026')
    baker.make('examens.InscriptionExamen', session=session, inscription=inscriptions['candidat'], statut='ADMIS')
    baker.make('examens.InscriptionExamen', session=session, inscription=inscriptions['candidate'], statut='ABSENT')
    return {'etab': etab, 'annee': annee, 'session': session, 'inscriptions': inscriptions}


@pytest.mark.django_db
class TestCalculs:
    def test_rentree(self, ecole):
        rentree = statistiques_rentree(ecole['annee'])
        sixieme, troisieme = rentree['lignes']
        assert (sixieme['niveau'], sixieme['divisions'], sixieme['effectif'], sixieme['redoublants'], sixieme['nouveaux']) == (
            '6ème', 2, {'M': 1, 'F': 2, 'T': 3}, {'M': 1, 'F': 0, 'T': 1}, {'M': 0, 'F': 2, 'T': 2})
        assert (troisieme['niveau'], troisieme['divisions'], troisieme['effectif']) == ('3ème', 1, {'M': 1, 'F': 1, 'T': 2})
        assert rentree['total']['effectif'] == {'M': 2, 'F': 3, 'T': 5} and rentree['niveaux'] == ['6ème', '3ème']
        ages = {ligne['age']: [c['T'] for c in ligne['colonnes']] for ligne in rentree['ages']}
        assert ages == {'11 ans': [2, 0], '12 ans': [1, 0], '14 ans': [0, 1], '15 ans': [0, 1]}
        # Pied du tableau des âges = effectif par niveau, et la somme des lignes d'âge retombe sur l'effectif total
        assert rentree['ages_total'] == {'colonnes': [sixieme['effectif'], troisieme['effectif']], 'total': {'M': 2, 'F': 3, 'T': 5}}
        assert {cle: sum(ligne['total'][cle] for ligne in rentree['ages']) for cle in 'MFT'} == rentree['total']['effectif']

    def test_meme_niveau_dans_deux_cycles(self, ecole):
        """Un niveau présent dans deux cycles (ex. filière franco-arabe) : colonnes distinctes, totaux non doublés."""
        etab, annee = ecole['etab'], ecole['annee']
        franco_arabe = baker.make('parametres.Cycle', etablissement=etab, nom='Franco-arabe', ordre=2)
        classe = baker.make('parametres.Classe', etablissement=etab, cycle=franco_arabe, niveau='6ème', nom='6ème FA')
        baker.make('inscriptions.Inscription', classe=classe, annee_scolaire=annee, statut='AFFECTE',
                   eleve=baker.make('inscriptions.Eleve', genre='F', date_naissance=datetime.date(2014, 2, 2)))
        rentree = statistiques_rentree(annee)
        assert rentree['niveaux'] == ['6ème (Post-primaire)', '3ème', '6ème (Franco-arabe)']
        onze_ans = rentree['ages'][0]
        assert onze_ans['age'] == '11 ans' and [(c['M'], c['F']) for c in onze_ans['colonnes']] == [(0, 2), (0, 0), (0, 1)]
        assert onze_ans['total'] == {'M': 0, 'F': 3, 'T': 3}
        assert rentree['ages_total']['total'] == rentree['total']['effectif'] == {'M': 2, 'F': 4, 'T': 6}

    def test_fin_annee(self, ecole):
        fin = statistiques_fin_annee(ecole['annee'])
        sixieme, troisieme = fin['lignes']
        assert (sixieme['inscrits']['T'], sixieme['abandons'], sixieme['fin_annee']['T']) == (3, {'M': 0, 'F': 1, 'T': 1}, 2)
        assert (sixieme['redoublement']['M'], sixieme['sans_decision']) == (1, {'M': 0, 'F': 1, 'T': 1})
        assert (troisieme['passage'], troisieme['sans_decision']['T']) == ({'M': 1, 'F': 1, 'T': 2}, 0)  # rachat compté en passage
        examen, = fin['examens']
        assert (examen['inscrits'], examen['presents'], examen['admis'], examen['taux']) == (
            {'M': 1, 'F': 1, 'T': 2}, {'M': 1, 'F': 0, 'T': 1}, {'M': 1, 'F': 0, 'T': 1}, {'M': 100.0, 'F': None, 'T': 100.0})

    def test_annee_sans_inscription(self, ecole):
        vide = baker.make('parametres.AnneeScolaire', etablissement=ecole['etab'], libelle='2030-2031',
                          date_debut=datetime.date(2030, 10, 1), date_fin=datetime.date(2031, 6, 30))
        assert statistiques_rentree(vide)['lignes'] == [] and statistiques_fin_annee(vide)['examens'] == []


@pytest.mark.django_db
class TestEcranEtExport:
    def test_ecran(self, client, ecole):
        client.force_login(baker.make('accounts.User', role='SECRETAIRE', etablissement=ecole['etab'], must_change_password=False))
        html = client.get(reverse('documents:statistiques_officielles')).content.decode()
        for attendu in ('Rentrée 2025-2026', '6ème', '3ème', 'Nouveaux entrants', '11 ans', '15 ans',
                        "Fin d'année 2025-2026", 'BEPC 2026', 'Exporter (Excel)'):
            assert attendu in html, attendu
        assert '5 élèves' in html and '4 présents' in html

    def test_disposition(self, client, ecole):
        """Tableaux au style maison (classe .table définie) et défilants, groupes G/F/T séparés, totaux en pied."""
        from pathlib import Path
        client.force_login(baker.make('accounts.User', role='SECRETAIRE', etablissement=ecole['etab'], must_change_password=False))
        html = client.get(reverse('documents:statistiques_officielles')).content.decode()
        assert html.count('<div class="table-wrapper table-scroll">\n    <table class="table table-stats">') == 4
        assert html.count('<th colspan="3" class="col-groupe">') == 16 and html.count('<td class="td-right col-groupe">') >= 20
        assert html.count('<tfoot>') == 3 and '<td colspan="3">Total</td>' in html and html.count('<td>Total</td>') == 2
        css = (Path(__file__).resolve().parents[1] / 'static/css/yelen.css').read_text(encoding='utf-8')
        assert '.table th' in css and '.table-stats .col-groupe' in css and '.table-stats tfoot td' in css

    def test_export_excel(self, client, ecole):
        client.force_login(baker.make('accounts.User', role='DIRECTEUR', etablissement=ecole['etab'], must_change_password=False))
        reponse = client.get(reverse('documents:statistiques_officielles'), {'annee_id': ecole['annee'].pk, 'format': 'xlsx'})
        assert reponse['Content-Disposition'].endswith('Statistiques_officielles_2025-2026.xlsx"')
        classeur = openpyxl.load_workbook(BytesIO(reponse.content))
        assert classeur.sheetnames == ['Rentrée', 'Âges', "Fin d'année"]
        rentree = list(classeur['Rentrée'].iter_rows(values_only=True))
        assert rentree[3][:6] == ('Cycle', 'Niveau', 'Divisions', 'Garçons', 'Filles', 'Total')
        assert rentree[4] == ('Post-primaire', '6ème', 2, 1, 2, 3, 1, 0, 1, 0, 2, 2)
        assert rentree[6][:6] == ('TOTAL', None, 3, 2, 3, 5)
        ages = list(classeur['Âges'].iter_rows(values_only=True))
        assert ages[3] == ('Âge', '6ème G', '6ème F', '3ème G', '3ème F', 'Garçons', 'Filles', 'Total')
        assert ages[4] == ('11 ans', 0, 2, 0, 0, 0, 2, 2)
        assert ages[8] == ('TOTAL', 1, 2, 1, 1, 2, 3, 5)
        fin = list(classeur["Fin d'année"].iter_rows(values_only=True))
        assert fin[4][:11] == ('Post-primaire', '6ème', 1, 2, 3, 0, 1, 1, 1, 1, 2)
        assert [ligne for ligne in fin if ligne[0] == 'BEPC'][0][:14] == ('BEPC', 'BEPC 2026', 1, 1, 2, 1, 0, 1, 1, 0, 1, 100.0, None, 100.0)

    def test_droits_et_isolation(self, client, ecole):
        client.force_login(baker.make('accounts.User', role='ENSEIGNANT', etablissement=ecole['etab'], must_change_password=False))
        assert client.get(reverse('documents:statistiques_officielles')).status_code in (302, 403)
        autre = baker.make('accounts.User', role='DIRECTEUR', etablissement=baker.make('etablissements.Etablissement'), must_change_password=False)
        client.force_login(autre)
        reponse = client.get(reverse('documents:statistiques_officielles'), {'annee_id': ecole['annee'].pk})
        assert reponse.status_code == 200 and 'Aucune année scolaire' in reponse.content.decode()
