"""
Moteur commun des assistants d'import (élèves, notes) : lecture d'un fichier
Excel (.xlsx) ou CSV tel qu'il existe dans l'école, détection de la ligne
d'entête, correspondance automatique des colonnes et conversions des valeurs.

Aucune dépendance aux vues : les assistants gardent le fichier téléversé dans
``MEDIA_ROOT/imports/`` entre deux étapes et le relisent à chaque étape (un
fichier de 2 500 lignes se lit en moins d'une seconde), ce qui évite de
stocker les données dans la session.
"""
import csv
import re
import unicodedata
import uuid
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.conf import settings

EXTENSIONS = ('.xlsx', '.xlsm', '.csv')
TAILLE_MAX = 10 * 1024 * 1024
LIGNES_MAX = 10_000
COLONNES_MAX = 80
LIGNES_ENTETE_MAX = 20
DOSSIER = 'imports'
ERREURS_MAX = 300  # erreurs conservées en session pour le rapport

_MOTS_VIDES = frozenset({'de', 'du', 'des', 'la', 'le', 'les', 'l', 'd', 'et', 'en', 'a', 'au', 'aux'})


# ── Normalisation ─────────────────────────────────────────────────────────

def texte(valeur) -> str:
    """Valeur de cellule → texte épuré ('' si vide) ; 70123456.0 → '70123456'."""
    if valeur is None:
        return ''
    if isinstance(valeur, float) and valeur.is_integer():
        valeur = int(valeur)
    if isinstance(valeur, datetime):
        return valeur.strftime('%d/%m/%Y')
    if isinstance(valeur, date):
        return valeur.strftime('%d/%m/%Y')
    return str(valeur).strip()


def nombre(valeur) -> str:
    """Decimal → texte sans zéros inutiles : 10.00 → « 10 », 12.50 → « 12.5 »."""
    brut = f'{Decimal(valeur):f}'
    return brut.rstrip('0').rstrip('.') if '.' in brut else brut


def normaliser(valeur) -> str:
    """Minuscules, sans accents ni ponctuation, espaces simples : « Date de naissance : » → « date de naissance »."""
    brut = unicodedata.normalize('NFKD', texte(valeur))
    brut = ''.join(c for c in brut if not unicodedata.combining(c)).lower()
    return ' '.join(re.sub(r'[^a-z0-9]+', ' ', brut).split())


def initiales(valeur) -> str:
    """« Sciences de la vie et de la terre » → « svt »."""
    return ''.join(m[0] for m in normaliser(valeur).split() if m not in _MOTS_VIDES)


def lettre_colonne(index: int) -> str:
    """0 → A, 25 → Z, 26 → AA."""
    lettres = ''
    index += 1
    while index:
        index, reste = divmod(index - 1, 26)
        lettres = chr(65 + reste) + lettres
    return lettres


# ── Fichier téléversé ─────────────────────────────────────────────────────

class ErreurImport(Exception):
    """Erreur à afficher telle quelle à l'utilisateur."""


def _dossier() -> Path:
    dossier = Path(settings.MEDIA_ROOT) / DOSSIER
    dossier.mkdir(parents=True, exist_ok=True)
    return dossier


def enregistrer_fichier(fichier) -> str:
    """Contrôle et enregistre le fichier téléversé ; retourne son nom interne (uuid + extension)."""
    extension = Path(fichier.name).suffix.lower()
    if extension not in EXTENSIONS:
        raise ErreurImport("Format non pris en charge : utilisez un fichier Excel (.xlsx) ou CSV. "
                           "Un ancien fichier .xls s'enregistre en .xlsx depuis Excel (Fichier → Enregistrer sous).")
    if fichier.size > TAILLE_MAX:
        raise ErreurImport("Fichier trop volumineux (10 Mo maximum).")
    dossier = _dossier()
    limite = datetime.now() - timedelta(days=1)
    for ancien in dossier.iterdir():
        if ancien.is_file() and datetime.fromtimestamp(ancien.stat().st_mtime) < limite:
            ancien.unlink(missing_ok=True)
    nom = f'{uuid.uuid4().hex}{extension}'
    with open(dossier / nom, 'wb') as sortie:
        for morceau in fichier.chunks():
            sortie.write(morceau)
    return nom


def chemin_fichier(nom: str) -> Path:
    """Chemin d'un fichier enregistré ; ErreurImport s'il a disparu (purge, redémarrage)."""
    if not nom or Path(nom).name != nom:
        raise ErreurImport("Fichier introuvable : recommencez l'import.")
    chemin = _dossier() / nom
    if not chemin.is_file():
        raise ErreurImport("Le fichier téléversé n'est plus disponible : recommencez l'import.")
    return chemin


def supprimer_fichier(nom: str) -> None:
    try:
        chemin_fichier(nom).unlink(missing_ok=True)
    except ErreurImport:
        pass


# ── Lecture ───────────────────────────────────────────────────────────────

@dataclass
class Feuille:
    nom: str
    entetes: list          # libellés de colonnes (jamais vides : « Colonne D » à défaut)
    lignes: list           # tuples de valeurs, longueur = len(entetes)
    numeros: list          # numéro de ligne dans le fichier, pour chaque ligne de données
    ligne_entete: int      # numéro de la ligne d'entête dans le fichier


def noms_feuilles(nom: str) -> list:
    chemin = chemin_fichier(nom)
    if chemin.suffix == '.csv':
        return ['CSV']
    import openpyxl
    try:
        classeur = openpyxl.load_workbook(chemin, read_only=True)
    except Exception as exc:
        raise ErreurImport(f"Fichier Excel illisible : {exc}") from exc
    try:
        return list(classeur.sheetnames)
    finally:
        classeur.close()


def _lignes_brutes(chemin: Path, feuille: str):
    if chemin.suffix == '.csv':
        octets = chemin.read_bytes()
        try:
            contenu = octets.decode('utf-8-sig')
        except UnicodeDecodeError:
            contenu = octets.decode('cp1252', errors='replace')
        try:
            dialecte = csv.Sniffer().sniff(contenu[:4096], delimiters=';,\t')
        except csv.Error:
            dialecte = csv.excel
            dialecte.delimiter = ';'
        for ligne in csv.reader(contenu.splitlines(), dialecte):
            yield ligne
        return
    import openpyxl
    try:
        classeur = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    except Exception as exc:
        raise ErreurImport(f"Fichier Excel illisible : {exc}") from exc
    try:
        if feuille not in classeur.sheetnames:
            feuille = classeur.sheetnames[0]
        for ligne in classeur[feuille].iter_rows(values_only=True, max_col=COLONNES_MAX):
            yield ligne
    finally:
        classeur.close()


def lire_feuille(nom: str, feuille: str = '', synonymes=()) -> Feuille:
    """Lit la feuille demandée (la première à défaut) et repère la ligne d'entête.

    L'entête est, parmi les premières lignes, celle qui contient le plus de libellés
    connus (``synonymes`` normalisés) ; à défaut la première ligne d'au moins deux
    cellules textuelles. Les lignes de titre qui précèdent sont ignorées, comme les
    lignes entièrement vides.
    """
    chemin = chemin_fichier(nom)
    feuille = feuille or (noms_feuilles(nom)[0])
    connus = set(synonymes)
    brutes = []
    for index, ligne in enumerate(_lignes_brutes(chemin, feuille)):
        if index >= LIGNES_MAX + LIGNES_ENTETE_MAX:
            break
        brutes.append(tuple(ligne)[:COLONNES_MAX])

    meilleure, score_max = None, -1
    for index, ligne in enumerate(brutes[:LIGNES_ENTETE_MAX]):
        textes = [normaliser(v) for v in ligne if re.search('[a-z]', normaliser(v))]
        if len(textes) < 2:
            continue
        score = sum(1 for t in textes if t in connus)
        if score > score_max:
            meilleure, score_max = index, score
    if meilleure is None:
        raise ErreurImport("Aucune ligne d'entête trouvée : la première ligne du tableau doit contenir "
                           "les titres des colonnes (Nom, Prénom, …).")

    entete_brute = list(brutes[meilleure])
    while entete_brute and not texte(entete_brute[-1]):
        entete_brute.pop()
    entetes = [texte(v) or f'Colonne {lettre_colonne(i)}' for i, v in enumerate(entete_brute)]
    largeur = len(entetes)

    lignes, numeros = [], []
    for decalage, ligne in enumerate(brutes[meilleure + 1:], start=meilleure + 2):
        valeurs = tuple(ligne[:largeur]) + (None,) * max(0, largeur - len(ligne))
        if any(texte(v) for v in valeurs):
            lignes.append(valeurs)
            numeros.append(decalage)
    return Feuille(nom=feuille, entetes=entetes, lignes=lignes, numeros=numeros, ligne_entete=meilleure + 1)


# ── Correspondance des colonnes ───────────────────────────────────────────

@dataclass(frozen=True)
class Champ:
    code: str
    libelle: str
    synonymes: tuple
    obligatoire: bool = False
    aide: str = ''


def synonymes_de(champs) -> set:
    return {s for champ in champs for s in champ.synonymes}


def detecter_colonnes(entetes, champs) -> dict:
    """code de champ → index de colonne.

    1. égalité stricte entre l'entête normalisé et un synonyme ;
    2. sinon, pour les synonymes d'au moins deux mots, inclusion de leurs mots
       dans l'entête (« date de naissance de l'élève » ⊃ « date naissance »).
    Une colonne n'est attribuée qu'une fois.
    """
    normalises = [normaliser(e) for e in entetes]
    correspondance, prises = {}, set()
    for champ in champs:
        for index, entete in enumerate(normalises):
            if index not in prises and entete in champ.synonymes:
                correspondance[champ.code] = index
                prises.add(index)
                break
    for champ in champs:
        if champ.code in correspondance:
            continue
        for index, entete in enumerate(normalises):
            if index in prises:
                continue
            mots = entete.split()
            if any(len(s.split()) >= 2 and all(m in mots for m in s.split()) for s in champ.synonymes):
                correspondance[champ.code] = index
                prises.add(index)
                break
    return correspondance


# ── Conversions ───────────────────────────────────────────────────────────

_FORMATS_DATE = ('%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y', '%Y-%m-%d', '%d/%m/%y', '%d-%m-%y')


def en_date(valeur):
    """Cellule date Excel, numéro de série Excel ou texte jj/mm/aaaa → date ; None si vide ; ValueError sinon."""
    if valeur is None or texte(valeur) == '':
        return None
    if isinstance(valeur, datetime):
        return valeur.date()
    if isinstance(valeur, date):
        return valeur
    if isinstance(valeur, (int, float)) and 1000 < valeur < 80000:
        from openpyxl.utils.datetime import from_excel
        return from_excel(valeur).date()
    brut = re.split(r'[ T]', texte(valeur))[0]  # « 12/05/2010 00:00:00 » → « 12/05/2010 »
    for fmt in _FORMATS_DATE:
        try:
            return datetime.strptime(brut, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"date invalide « {texte(valeur)} » (attendu jj/mm/aaaa)")


_GENRES = {
    'M': ('m', 'masculin', 'garcon', 'g', 'h', 'homme', 'male', 'mr', 'boy'),
    'F': ('f', 'feminin', 'fille', 'femme', 'female', 'mme', 'mlle', 'girl'),
}


def en_genre(valeur):
    """M / F ; None si vide ; ValueError sinon."""
    brut = normaliser(valeur)
    if not brut:
        return None
    for code, formes in _GENRES.items():
        if brut in formes:
            return code
    raise ValueError(f"genre invalide « {texte(valeur)} » (attendu M ou F)")


_OUI = ('oui', 'o', 'x', '1', 'vrai', 'true', 'yes', 'y', 'r', 'red', 'redoublant', 'redoublante', 'doublant')
_NON = ('non', 'n', '0', 'faux', 'false', 'no', 'nouveau', 'nouvelle', 'passant', 'passante', '')


def en_booleen(valeur) -> bool:
    brut = normaliser(valeur)
    if brut in _OUI:
        return True
    if brut in _NON:
        return False
    raise ValueError(f"valeur invalide « {texte(valeur)} » (attendu Oui ou Non)")


def en_note(valeur, bareme):
    """(valeur Decimal, absent) ; None si la cellule est vide ; ValueError si invalide ou hors barème."""
    brut = texte(valeur)
    if brut == '' or brut in ('-', '—'):
        return None
    if normaliser(brut) in ('abs', 'absent', 'absente', 'a'):
        return Decimal('0.00'), True
    try:
        note = Decimal(brut.replace(',', '.').replace(' ', ''))
    except InvalidOperation:
        raise ValueError(f"note invalide « {brut} »") from None
    if note < 0 or note > bareme:
        raise ValueError(f"note {brut} hors barème (0 à {nombre(bareme)})")
    return note.quantize(Decimal('0.01')), False


def en_telephone(valeur) -> str:
    brut = texte(valeur)
    if not brut:
        return ''
    chiffres = re.sub(r'[^0-9+]', '', brut)
    return chiffres[:20]


def separer_nom_prenom(valeur):
    """« OUEDRAOGO Abdoul Kader » → ('OUEDRAOGO', 'Abdoul Kader').

    Les mots en MAJUSCULES qui ouvrent la chaîne forment le nom ; à défaut
    d'indice de casse, le premier mot est le nom et le reste les prénoms.
    """
    mots = texte(valeur).split()
    if not mots:
        return '', ''
    if len(mots) == 1:
        return mots[0], ''
    majuscules = 0
    for mot in mots:
        if mot.isupper() and len(mot) > 1:
            majuscules += 1
        else:
            break
    if 0 < majuscules < len(mots):
        return ' '.join(mots[:majuscules]), ' '.join(mots[majuscules:])
    return mots[0], ' '.join(mots[1:])


def cle_personne(nom, prenom) -> str:
    return f'{normaliser(nom)}|{normaliser(prenom)}'


def cle_nom_complet(nom, prenom) -> str:
    return normaliser(f'{nom} {prenom}')
