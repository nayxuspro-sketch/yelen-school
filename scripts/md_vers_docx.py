"""Convertit un guide Markdown du dossier docs/ en document Word (.docx).

Usage :
    python scripts/md_vers_docx.py docs/GUIDE_UTILISATION_YELEN_SCHOOL.md [sortie.docx]

Pris en charge : en-tête YAML (fiche d'identité), titres (# à #####), listes à puces
imbriquées, listes numérotées, tableaux, blocs de code, citations, images (chemins
relatifs au fichier .md), gras **x**, code `x`, liens [texte](url). La table des
matières manuelle du Markdown est remplacée par un champ TOC Word (mis à jour à
l'ouverture du document). Dépendance : python-docx (requirements/dev.txt).
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|!?\[[^\]]+\]\([^)]+\)|(?<![\w*])\*(?!\s)[^*\n]+?(?<!\s)\*(?![\w*]))")
CODE_FONT = "Consolas"


def _shade(element, fill):
    """Fond coloré (hex sans #) sur un paragraphe ou une cellule."""
    pr = element.get_or_add_pPr() if hasattr(element, "get_or_add_pPr") else element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _field(paragraph, instruction):
    """Insère un champ Word (TOC, PAGE…) dans un paragraphe."""
    run = paragraph.add_run()
    for tag, attrs, text in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, instruction),
        ("w:fldChar", {"w:fldCharType": "separate"}, None),
        ("w:t", {}, "Table des matières : clic droit → « Mettre à jour les champs »."),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        if text:
            el.text = text
        run._r.append(el)


def add_inline(paragraph, text, base_size=None):
    """Ajoute `text` au paragraphe en interprétant gras, code et liens."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt((base_size or 10.5) - 1)
            run.font.color.rgb = RGBColor(0x9C, 0x27, 0x06)
        elif part.startswith("[") or part.startswith("!["):
            m = re.match(r"!?\[([^\]]+)\]\(([^)]+)\)", part)
            label, url = m.group(1), m.group(2)
            run = paragraph.add_run(label)
            if url.startswith("http"):
                paragraph.add_run(f" ({url})").font.size = Pt((base_size or 10.5) - 1.5)
        else:
            run = paragraph.add_run(part)
        if base_size and run.font.size is None:
            run.font.size = Pt(base_size)


class Convertisseur:
    def __init__(self, md_path: Path):
        self.md_path = md_path
        self.doc = Document()
        self._mise_en_page()
        self.lines = md_path.read_text(encoding="utf-8").splitlines()
        self.meta = {}
        self.titre_pose = False

    # ----- préparation du document -------------------------------------------------
    def _mise_en_page(self):
        doc = self.doc
        section = doc.sections[0]
        section.page_height, section.page_width = Cm(29.7), Cm(21.0)
        for side in ("left_margin", "right_margin"):
            setattr(section, side, Cm(2.0))
        section.top_margin = section.bottom_margin = Cm(1.8)
        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(4)
        for name, size, color in (("Heading 1", 18, "1F3864"), ("Heading 2", 14, "2E5597"),
                                  ("Heading 3", 12, "2E5597"), ("Heading 4", 11, "404040")):
            st = doc.styles[name]
            st.font.name = "Calibri"
            st.font.size = Pt(size)
            st.font.color.rgb = RGBColor.from_string(color)
        # Mise à jour automatique des champs (table des matières) à l'ouverture
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        doc.settings.element.append(upd)
        # Pied de page : nom du guide + numéro de page
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("YELEN SCHOOL — Guide d'utilisation — page ").font.size = Pt(8)
        _field(footer, "PAGE")

    # ----- blocs ------------------------------------------------------------------
    def convertir(self, sortie: Path):
        i, n = 0, len(self.lines)
        if n and self.lines[0].strip() == "---":
            i = self._front_matter(1)
        dans_toc = False
        while i < n:
            line = self.lines[i]
            s = line.strip()
            if s.startswith("```"):
                i = self._bloc_code(i + 1)
                continue
            if s.startswith("|") and i + 1 < n and re.match(r"^\|?\s*:?-{3,}", self.lines[i + 1].strip()):
                i = self._tableau(i)
                continue
            m = re.match(r"^(#{1,6})\s+(.*)$", s)
            if m:
                niveau, texte = len(m.group(1)), m.group(2).strip()
                dans_toc = texte.upper().startswith("TABLE DES MATI")
                if dans_toc:
                    self.doc.add_heading("Table des matières", level=1)
                    _field(self.doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u')
                    self.doc.add_page_break()
                else:
                    self._titre(niveau, texte)
                i += 1
                continue
            if dans_toc or not s or s in ("---", "***"):
                i += 1
                continue
            m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", s)
            if m:
                self._image(m.group(1), m.group(2))
                i += 1
                continue
            if s.startswith(">"):
                i = self._citation(i)
                continue
            m = re.match(r"^(\s*)([-*+])\s+(.*)$", line)
            if m:
                niveau = min(len(m.group(1).expandtabs(4)) // 2, 2)
                style = "List Bullet" if niveau == 0 else f"List Bullet {niveau + 1}"
                texte, i = self._suite_element(m.group(3), i + 1)
                add_inline(self.doc.add_paragraph(style=style), texte)
                continue
            m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
            if m:
                p = self.doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.75 + 0.5 * (len(m.group(1)) // 2))
                p.paragraph_format.first_line_indent = Cm(-0.6)
                texte, i = self._suite_element(f"{m.group(2)}. {m.group(3)}", i + 1)
                add_inline(p, texte)
                continue
            # paragraphe : fusion des lignes consécutives
            buf = [s]
            i += 1
            while i < n and self.lines[i].strip() and not re.match(
                    r"^(#{1,6}\s|```|\||>|\s*[-*+]\s|\s*\d+[.)]\s|!\[)", self.lines[i]):
                buf.append(self.lines[i].strip())
                i += 1
            add_inline(self.doc.add_paragraph(), " ".join(buf))
        self.doc.save(sortie)

    def _suite_element(self, texte, i):
        """Lignes de continuation (indentées) d'un élément de liste."""
        while i < len(self.lines) and self.lines[i].strip() and self.lines[i][:1] in (" ", "\t") and not re.match(
                r"^\s*([-*+]\s|\d+[.)]\s|```|\||>|!\[)", self.lines[i]):
            texte += " " + self.lines[i].strip()
            i += 1
        return texte, i

    def _front_matter(self, i):
        while i < len(self.lines) and self.lines[i].strip() != "---":
            if ":" in self.lines[i]:
                k, v = self.lines[i].split(":", 1)
                self.meta[k.strip()] = v.strip()
            i += 1
        return i + 1

    def _titre(self, niveau, texte):
        if niveau == 1 and not self.titre_pose:
            self.titre_pose = True
            p = self.doc.add_paragraph(style="Title")
            add_inline(p, texte)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if self.meta:
                table = self.doc.add_table(rows=0, cols=2)
                table.style = "Table Grid"
                for k in ("titre", "version_logiciel", "version_guide", "date_mise_a_jour", "redige_par"):
                    if k in self.meta:
                        cells = table.add_row().cells
                        cells[0].text = k.replace("_", " ").capitalize()
                        cells[0].paragraphs[0].runs[0].bold = True
                        cells[1].text = self.meta[k]
                        _shade(cells[0]._tc, "DEEAF6")
                self.doc.add_paragraph()
            return
        if niveau == 1:
            niveau = 2
        texte = texte.strip("*").strip()
        self.doc.add_heading(texte, level=min(niveau - 1, 4))

    def _bloc_code(self, i):
        first = True
        while i < len(self.lines) and not self.lines[i].strip().startswith("```"):
            p = self.doc.add_paragraph(style="No Spacing")
            p.paragraph_format.left_indent = Cm(0.5)
            if first:
                p.paragraph_format.space_before = Pt(4)
                first = False
            run = p.add_run(self.lines[i].rstrip() or " ")
            run.font.name = CODE_FONT
            run.font.size = Pt(8.5)
            _shade(p._p, "F2F2F2")
            i += 1
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return i + 1

    def _tableau(self, i):
        rows = []
        while i < len(self.lines) and self.lines[i].strip().startswith("|"):
            s = self.lines[i].strip().strip("|")
            if not re.match(r"^\s*:?-{3,}", s):
                rows.append([c.strip() for c in re.split(r"(?<!\\)\|", s)])
            i += 1
        if not rows:
            return i
        ncols = max(len(r) for r in rows)
        table = self.doc.add_table(rows=0, cols=ncols)
        table.style = "Table Grid"
        for r_idx, row in enumerate(rows):
            cells = table.add_row().cells
            for c_idx in range(ncols):
                texte = row[c_idx] if c_idx < len(row) else ""
                p = cells[c_idx].paragraphs[0]
                add_inline(p, texte.replace("\\|", "|"), base_size=9)
                if r_idx == 0:
                    for run in p.runs:
                        run.bold = True
                    _shade(cells[c_idx]._tc, "DEEAF6")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return i

    def _citation(self, i):
        buf = []
        while i < len(self.lines) and self.lines[i].strip().startswith(">"):
            buf.append(self.lines[i].strip().lstrip(">").strip())
            i += 1
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        _shade(p._p, "FFF8E1")
        add_inline(p, " ".join(b for b in buf if b))
        return i

    def _image(self, alt, chemin):
        fichier = (self.md_path.parent / chemin).resolve()
        if fichier.exists():
            self.doc.add_picture(str(fichier), width=Cm(16))
            self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if alt:
                cap = self.doc.add_paragraph(alt, style="Caption")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            self.doc.add_paragraph(f"[Illustration : {alt or chemin}]").runs[0].italic = True


def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 1
    source = Path(argv[1])
    sortie = Path(argv[2]) if len(argv) > 2 else source.with_suffix(".docx")
    Convertisseur(source).convertir(sortie)
    print(f"{sortie} ({sortie.stat().st_size // 1024} Ko)")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
