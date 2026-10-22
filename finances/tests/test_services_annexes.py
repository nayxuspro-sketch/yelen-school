"""Services annexes facturés (cantine, transport, internat) : calcul, situation financière, vues."""
from datetime import date
from decimal import Decimal

import pytest
from django.urls import reverse
from model_bakery import baker
from openpyxl import load_workbook

from finances.models import AbonnementService, FormuleService, ServiceAnnexe
from finances.views.commun import _calcul_situation_financiere, _payment_overage_errors, _situations_financieres_en_lot
from parametres.models import RubriquePaiement


@pytest.fixture
def ecole():
    etab = baker.make('etablissements.Etablissement')
    annee = baker.make('parametres.AnneeScolaire', etablissement=etab, est_courante=True,
                       date_debut=date(2026, 10, 1), date_fin=date(2027, 6, 30))
    classe = baker.make('parametres.Classe', etablissement=etab, nom='6e A', niveau='6e', actif=True)
    rubrique_cantine = baker.make('parametres.RubriquePaiement', etablissement=etab, nom='Cantine', code='CANT', actif=True)
    cantine = ServiceAnnexe.objects.create(etablissement=etab, type_service='CANTINE', nom='Cantine du midi',
                                           rubrique=rubrique_cantine, periodicite='MENSUEL')
    formule = FormuleService.objects.create(service=cantine, libelle='Déjeuner 5 j/7', montant=5000)
    inscription = baker.make('inscriptions.Inscription', classe=classe, annee_scolaire=annee, statut_eleve=None,
                             eleve__nom='OUEDRAOGO', eleve__prenom='Awa', eleve__date_naissance=date(2014, 5, 3))
    return {'etab': etab, 'annee': annee, 'classe': classe, 'cantine': cantine, 'formule': formule,
            'rubrique': rubrique_cantine, 'inscription': inscription}


@pytest.fixture
def comptable(client, ecole):
    user = baker.make('accounts.User', etablissement=ecole['etab'], role='COMPTABLE')
    client.force_login(user)
    return client


def _abonner(ecole, debut=date(2026, 10, 1), fin=date(2027, 6, 30), inscription=None):
    abonnement = AbonnementService(inscription=inscription or ecole['inscription'], formule=ecole['formule'],
                                   date_debut=debut, date_fin=fin)
    abonnement.save()
    return abonnement


# ── Calcul ───────────────────────────────────────────────────────────────

@pytest.mark.django_db
@pytest.mark.parametrize('periodicite, debut, fin, attendu', [
    ('MENSUEL', date(2026, 10, 1), date(2027, 6, 30), 9),
    ('MENSUEL', date(2026, 10, 15), date(2026, 10, 20), 1),
    ('MENSUEL', date(2026, 12, 20), date(2027, 1, 5), 2),
    ('TRIMESTRIEL', date(2026, 10, 1), date(2027, 6, 30), 3),
    ('TRIMESTRIEL', date(2026, 10, 1), date(2026, 11, 30), 1),
    ('TRIMESTRIEL', date(2026, 10, 1), date(2027, 1, 31), 2),
    ('ANNUEL', date(2026, 10, 1), date(2027, 6, 30), 1),
])
def test_nb_periodes(ecole, periodicite, debut, fin, attendu):
    ecole['cantine'].periodicite = periodicite
    assert ecole['cantine'].nb_periodes(debut, fin) == attendu


@pytest.mark.django_db
def test_abonnement_fige_le_montant(ecole):
    abonnement = _abonner(ecole)
    assert (abonnement.nb_periodes, abonnement.montant_periode, abonnement.montant_total) == (9, 5000, 45000)
    assert abonnement.periodes_libelle == "9 mois du 01/10/2026 au 30/06/2027"
    ecole['formule'].montant = 7000
    ecole['formule'].save()
    abonnement.refresh_from_db()
    assert abonnement.montant_total == 45000  # le tarif figé ne bouge pas avec la formule


# ── Situation financière ────────────────────────────────────────────────

@pytest.mark.django_db
def test_situation_sans_statut_eleve_compte_les_services(ecole):
    """Sans tarif de scolarité applicable, l'abonnement est tout de même dû (unitaire et en lot)."""
    _abonner(ecole)
    sit = _calcul_situation_financiere(ecole['inscription'])
    assert sit['total_du'] == 45000 and sit['reste_a_payer'] == 45000
    assert [(f.libelle, f.montant) for f in sit['frais_details']] == [('Cantine du midi — Déjeuner 5 j/7', 45000)]
    lot = _situations_financieres_en_lot([ecole['inscription']])
    assert lot[ecole['inscription'].pk]['total_du'] == 45000


@pytest.mark.django_db
def test_situation_tarifs_plus_services_moins_paiement(ecole):
    statut = baker.make('parametres.StatutEleve')
    rubrique_scol = baker.make('parametres.RubriquePaiement', etablissement=ecole['etab'], code='SCOL', actif=True)
    baker.make('parametres.TarifScolarite', etablissement=ecole['etab'], classe=ecole['classe'],
               annee_scolaire=ecole['annee'], statut_eleve=statut, rubrique=rubrique_scol, montant=50000, actif=True)
    ecole['inscription'].statut_eleve = statut
    ecole['inscription'].save()
    _abonner(ecole)
    baker.make('finances.Paiement', inscription=ecole['inscription'], rubrique=ecole['rubrique'], montant=20000)

    sit = _calcul_situation_financiere(ecole['inscription'])
    assert (sit['total_du'], sit['total_paye'], sit['reste_a_payer']) == (95000, 20000, 75000)
    assert _situations_financieres_en_lot([ecole['inscription']])[ecole['inscription'].pk]['reste_a_payer'] == 75000


@pytest.mark.django_db
def test_abonnement_annule_ne_compte_plus(ecole):
    abonnement = _abonner(ecole)
    abonnement.actif = False
    abonnement.save()
    assert _calcul_situation_financiere(ecole['inscription'])['total_du'] == 0


@pytest.mark.django_db
def test_plafond_paiement_sur_rubrique_service(ecole):
    """Sans tarif, le montant de l'abonnement plafonne les versements sur sa rubrique."""
    _abonner(ecole)
    assert _payment_overage_errors(ecole['inscription'], [(ecole['rubrique'], Decimal('45000'), None)]) == []
    erreurs = _payment_overage_errors(ecole['inscription'], [(ecole['rubrique'], Decimal('45001'), None)])
    assert len(erreurs) == 1 and 'dépasse' in erreurs[0]


@pytest.mark.django_db
def test_api_rubriques_expose_le_service(comptable, ecole):
    _abonner(ecole)
    data = comptable.get(reverse('finances:api_rubriques', kwargs={'inscription_id': ecole['inscription'].pk})).json()
    assert 'error' not in data
    assert [(r['code'], r['montant'], r['reste']) for r in data['rubriques']] == [('CANT', '45000', '45000')]
    assert data['situation']['total_du'] == '45000'


# ── Vues ────────────────────────────────────────────────────────────────

@pytest.mark.django_db
def test_service_create_cree_la_rubrique_dediee(comptable, ecole):
    reponse = comptable.post(reverse('finances:service_create'), {
        'type_service': 'TRANSPORT', 'nom': 'Transport scolaire', 'periodicite': 'TRIMESTRIEL', 'rubrique': '', 'actif': 'on',
    })
    assert reponse.status_code == 302
    service = ServiceAnnexe.objects.get(etablissement=ecole['etab'], type_service='TRANSPORT')
    assert (service.rubrique.code, service.rubrique.nom, service.rubrique.obligatoire) == ('TRANS', 'Transport scolaire', False)
    assert RubriquePaiement.objects.filter(etablissement=ecole['etab'], code='TRANS').count() == 1


@pytest.mark.django_db
def test_formule_create_refuse_montant_nul(comptable, ecole):
    url = reverse('finances:formule_create', kwargs={'service_id': ecole['cantine'].pk})
    assert comptable.post(url, {'libelle': 'Goûter', 'montant': '0'}).status_code == 200
    assert not FormuleService.objects.filter(libelle='Goûter').exists()
    assert comptable.post(url, {'libelle': 'Goûter', 'montant': '2 500', 'actif': 'on'}).status_code == 302
    assert FormuleService.objects.get(libelle='Goûter').montant == 2500


@pytest.mark.django_db
def test_abonnement_create_puis_doublon_refuse(comptable, ecole):
    url = reverse('finances:abonnement_create', kwargs={'inscription_id': ecole['inscription'].pk})
    donnees = {'formule': ecole['formule'].pk, 'date_debut': '2026-11-01', 'date_fin': '2027-06-30'}
    reponse = comptable.post(url, donnees)
    assert reponse.status_code == 302 and reponse.url == reverse('finances:situation_eleve', args=[ecole['inscription'].pk])
    abonnement = AbonnementService.objects.get(inscription=ecole['inscription'])
    assert (abonnement.nb_periodes, abonnement.montant_total) == (8, 40000)

    reponse = comptable.post(url, donnees)  # second abonnement au même service
    assert reponse.status_code == 200
    assert AbonnementService.objects.filter(inscription=ecole['inscription']).count() == 1
    assert 'déjà abonné' in reponse.content.decode()


@pytest.mark.django_db
def test_abonnement_edit_recalcule_et_annuler(comptable, ecole):
    abonnement = _abonner(ecole)
    reponse = comptable.post(reverse('finances:abonnement_edit', kwargs={'pk': abonnement.pk}), {
        'formule': ecole['formule'].pk, 'date_debut': '2027-01-01', 'date_fin': '2027-03-31', 'observation': 'Retour',
    })
    assert reponse.status_code == 302
    abonnement.refresh_from_db()
    assert (abonnement.nb_periodes, abonnement.montant_total, abonnement.observation) == (3, 15000, 'Retour')

    reponse = comptable.post(reverse('finances:abonnement_annuler', kwargs={'pk': abonnement.pk}))
    assert reponse.status_code == 302
    abonnement.refresh_from_db()
    assert abonnement.actif is False
    assert _calcul_situation_financiere(ecole['inscription'])['total_du'] == 0


@pytest.mark.django_db
def test_situation_eleve_affiche_les_services(comptable, ecole):
    _abonner(ecole)
    html = comptable.get(reverse('finances:situation_eleve', kwargs={'inscription_id': ecole['inscription'].pk})).content.decode()
    assert 'Cantine du midi — Déjeuner 5 j/7' in html
    assert '9 mois du 01/10/2026 au 30/06/2027' in html
    assert reverse('finances:abonnement_create', kwargs={'inscription_id': ecole['inscription'].pk}) in html


@pytest.mark.django_db
def test_abonnes_list_html_et_excel(comptable, ecole):
    _abonner(ecole)
    baker.make('finances.Paiement', inscription=ecole['inscription'], rubrique=ecole['rubrique'], montant=10000)
    url = reverse('finances:abonnes_list', kwargs={'pk': ecole['cantine'].pk})
    reponse = comptable.get(url)
    assert reponse.status_code == 200
    a = reponse.context['abonnements'][0]
    assert (a.verse, a.reste) == (10000, 35000)
    assert reponse.context['totaux'] == {'du': 45000, 'verse': 10000, 'reste': 35000}

    reponse = comptable.get(url + '?format=xlsx')
    assert reponse['Content-Type'].startswith('application/vnd.openxmlformats')
    import io
    ws = load_workbook(io.BytesIO(reponse.content)).active
    lignes = [[c.value for c in r] for r in ws.iter_rows()]
    assert ['Classe', 'Matricule', 'Nom', 'Prénom(s)', 'Formule', 'Du', 'Au', 'Périodes', 'Montant dû', 'Versé', 'Reste', 'Téléphone parents'] in lignes
    assert any(l[:3] == ['6e A', ecole['inscription'].eleve.matricule, 'OUEDRAOGO'] and l[8:11] == [45000, 10000, 35000] for l in lignes)


@pytest.mark.django_db
def test_abonnement_classe_en_lot_ignore_les_deja_abonnes(comptable, ecole):
    _abonner(ecole)  # déjà abonné
    autres = [baker.make('inscriptions.Inscription', classe=ecole['classe'], annee_scolaire=ecole['annee'],
                         eleve__date_naissance=date(2014, 1, 1)) for _ in range(2)]
    abandon = baker.make('inscriptions.Inscription', classe=ecole['classe'], annee_scolaire=ecole['annee'],
                         statut='ABANDON', eleve__date_naissance=date(2014, 1, 1))
    url = reverse('finances:abonnement_classe', kwargs={'pk': ecole['cantine'].pk})
    page = comptable.get(url + f"?classe={ecole['classe'].pk}")
    assert page.status_code == 200
    assert [i.pk for i in page.context['eleves']].count(abandon.pk) == 0
    assert page.context['nb_restants'] == 2

    reponse = comptable.post(url, {
        'classe': ecole['classe'].pk, 'valider': '1', 'formule': ecole['formule'].pk,
        'date_debut': '2026-10-01', 'date_fin': '2027-06-30',
        'inscriptions': [str(ecole['inscription'].pk)] + [str(i.pk) for i in autres],
    })
    assert reponse.status_code == 302
    assert AbonnementService.objects.filter(formule=ecole['formule'], actif=True).count() == 3
    assert AbonnementService.objects.filter(inscription=ecole['inscription']).count() == 1


@pytest.mark.django_db
def test_service_list_totaux_et_droits(client, comptable, ecole):
    _abonner(ecole)
    reponse = comptable.get(reverse('finances:service_list'))
    assert reponse.status_code == 200
    service = reponse.context['services'][0]
    assert (service.nb_abonnes, service.total_du) == (1, 45000)

    client.logout()
    client.force_login(baker.make('accounts.User', etablissement=ecole['etab'], role='ENSEIGNANT'))
    assert client.get(reverse('finances:service_list')).status_code in (302, 403)
    assert client.post(reverse('finances:abonnement_create', kwargs={'inscription_id': ecole['inscription'].pk}),
                       {'formule': ecole['formule'].pk}).status_code in (302, 403)
    assert AbonnementService.objects.count() == 1


@pytest.mark.django_db
def test_autre_etablissement_refuse(comptable, ecole):
    autre = baker.make('inscriptions.Inscription', eleve__date_naissance=date(2014, 1, 1))
    reponse = comptable.get(reverse('finances:abonnement_create', kwargs={'inscription_id': autre.pk}))
    assert reponse.status_code in (302, 403)


@pytest.mark.django_db
def test_redevables_inclut_un_abonne_sans_scolarite(comptable, ecole):
    """L'élève sans tarif de scolarité mais abonné à la cantine apparaît dans les redevables."""
    _abonner(ecole)
    reponse = comptable.get(reverse('finances:liste_redevables'))
    assert reponse.status_code == 200
    assert 'OUEDRAOGO' in reponse.content.decode()
