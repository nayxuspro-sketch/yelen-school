"""Bourses, types de bourse et échéanciers : calculs, validations, filtres et exports."""
import io
from datetime import date, timedelta
from decimal import Decimal
from types import SimpleNamespace

import pytest
from django.contrib.messages import get_messages
from django.urls import reverse
from django.utils import timezone
from model_bakery import baker
from openpyxl import load_workbook

from finances.models import BourseEleve, Echeancier, TypeBourse


def _messages(reponse):
    return [str(m) for m in get_messages(reponse.wsgi_request)]


@pytest.fixture
def ctx(client):
    etab = baker.make('etablissements.Etablissement', nom='LE SAVOIR')
    user = baker.make('accounts.User', is_superuser=True, role='SUPER_ADMIN', etablissement=etab,
                      must_change_password=False)
    client.force_login(user)
    annee = baker.make('parametres.AnneeScolaire', etablissement=etab, libelle='2026-2027', est_courante=True,
                       date_debut=date(2026, 10, 1), date_fin=date(2027, 6, 30))
    cycle = baker.make('parametres.Cycle', etablissement=etab, code='SEC')
    classe = baker.make('parametres.Classe', etablissement=etab, cycle=cycle, nom='6e A', actif=True)
    eleve = baker.make('inscriptions.Eleve', nom='OUEDRAOGO', prenom='Awa', matricule='M-001', photo='')
    inscription = baker.make('inscriptions.Inscription', eleve=eleve, classe=classe, annee_scolaire=annee,
                             statut='ACTIF')
    rubrique = baker.make('parametres.RubriquePaiement', etablissement=etab, nom='Scolarité', actif=True)
    pct = baker.make('finances.TypeBourse', etablissement=etab, nom='Excellence', code='EXC', source='ETABLISSEMENT',
                     type_reduction='POURCENTAGE', valeur_reduction=Decimal('25'), actif=True)
    fixe = baker.make('finances.TypeBourse', etablissement=etab, nom='Aide ONG', code='ONG1', source='ONG',
                      type_reduction='MONTANT_FIXE', valeur_reduction=Decimal('30000'), actif=True)
    return SimpleNamespace(client=client, user=user, etab=etab, annee=annee, classe=classe, inscription=inscription,
                           rubrique=rubrique, pct=pct, fixe=fixe)


def _bourse(ctx, type_bourse, montant, inscription=None, **extra):
    return baker.make('finances.BourseEleve', inscription=inscription or ctx.inscription, type_bourse=type_bourse,
                      montant_accorde=Decimal(montant), actif=True, is_active=True, **extra)


@pytest.mark.django_db
class TestCalculBourse:
    def test_pourcentage_montant_fixe_et_type_inconnu(self, ctx):
        url = reverse('finances:api_calculer_bourse')
        assert ctx.client.get(url, {'type_bourse': ctx.pct.pk, 'total_du': '100000'}).json()['montant'] == '25000'
        assert Decimal(ctx.client.get(url, {'type_bourse': ctx.fixe.pk, 'total_du': '50000'}).json()['montant']) == 30000
        assert Decimal(ctx.client.get(url, {'type_bourse': ctx.fixe.pk, 'total_du': '20000'}).json()['montant']) == 20000, \
            'un montant fixe ne dépasse jamais le total dû'
        assert ctx.client.get(url, {'type_bourse': ctx.pct.pk, 'total_du': 'abc'}).json()['montant'] == '0'
        autre = baker.make('finances.TypeBourse', type_reduction='POURCENTAGE', valeur_reduction=Decimal('50'), actif=True)
        assert ctx.client.get(url, {'type_bourse': autre.pk, 'total_du': '100000'}).json() == {'montant': '0'}, \
            'type d’un autre établissement'


@pytest.mark.django_db
class TestAttributionBourse:
    def test_attribution_valide(self, ctx):
        url = reverse('finances:bourse_create', kwargs={'inscription_id': ctx.inscription.pk})
        assert ctx.client.get(url).status_code == 200
        reponse = ctx.client.post(url, {
            'type_bourse': ctx.pct.pk, 'montant_accorde': '25000,50', 'rubrique': ctx.rubrique.pk,
            'date_attribution': '2026-10-15', 'reference_document': 'DEC-12', 'observation': 'Major de promotion',
        })
        assert reponse.status_code == 302
        assert reponse.url == reverse('finances:situation_eleve', kwargs={'inscription_id': ctx.inscription.pk})
        bourse = BourseEleve.objects.get()
        assert (bourse.montant_accorde, bourse.type_bourse, bourse.rubrique) == (Decimal('25000.50'), ctx.pct, ctx.rubrique)
        assert bourse.date_attribution == date(2026, 10, 15) and bourse.date_expiration is None
        assert bourse.attribue_par == ctx.user and bourse.actif

    @pytest.mark.parametrize('donnees, erreur', [
        ({'montant_accorde': '-5'}, 'Le montant doit être positif.'),
        ({'montant_accorde': 'abc'}, 'Montant invalide.'),
        ({'type_bourse': 'inconnu', 'montant_accorde': '1000'}, 'Type de bourse invalide.'),
    ])
    def test_attribution_refusee(self, ctx, donnees, erreur):
        url = reverse('finances:bourse_create', kwargs={'inscription_id': ctx.inscription.pk})
        reponse = ctx.client.post(url, {'type_bourse': ctx.pct.pk, **donnees})
        assert reponse.status_code == 200 and erreur in _messages(reponse)
        assert not BourseEleve.objects.exists()

    def test_type_autre_etablissement_refuse(self, ctx):
        autre = baker.make('finances.TypeBourse', type_reduction='POURCENTAGE', valeur_reduction=Decimal('50'), actif=True)
        url = reverse('finances:bourse_create', kwargs={'inscription_id': ctx.inscription.pk})
        reponse = ctx.client.post(url, {'type_bourse': autre.pk, 'montant_accorde': '1000'})
        assert 'Type de bourse invalide.' in _messages(reponse) and not BourseEleve.objects.exists()

    def test_modification(self, ctx):
        bourse = _bourse(ctx, ctx.pct, '10000')
        url = reverse('finances:bourse_edit', kwargs={'bourse_id': bourse.pk})
        assert ctx.client.get(url).context['bourse'] == bourse
        reponse = ctx.client.post(url, {'type_bourse': ctx.fixe.pk, 'montant_accorde': '12000',
                                        'date_attribution': '2026-11-01', 'date_expiration': '2027-06-30'})
        assert reponse.status_code == 302
        bourse.refresh_from_db()
        assert (bourse.type_bourse, bourse.montant_accorde) == (ctx.fixe, Decimal('12000'))
        assert bourse.date_expiration == date(2027, 6, 30)
        assert bourse.actif is False, 'case « actif » non cochée'
        assert bourse.updated_by == ctx.user

    def test_desactivation_exige_un_motif(self, ctx):
        bourse = _bourse(ctx, ctx.pct, '10000')
        url = reverse('finances:bourse_delete', kwargs={'bourse_id': bourse.pk})
        assert ctx.client.get(url).status_code == 405
        reponse = ctx.client.post(url, {'motif': ''})
        bourse.refresh_from_db()
        assert bourse.actif and 'Le motif de désactivation est obligatoire.' in _messages(reponse)
        ctx.client.post(url, {'motif': 'Erreur de saisie'})
        bourse.refresh_from_db()
        assert (bourse.actif, bourse.is_active) == (False, False), 'désactivée, jamais supprimée'
        assert BourseEleve.objects.filter(pk=bourse.pk).exists()


@pytest.mark.django_db
class TestBoursiers:
    @pytest.fixture
    def bourses(self, ctx):
        b1 = _bourse(ctx, ctx.pct, '25000')
        autre_classe = baker.make('parametres.Classe', etablissement=ctx.etab, cycle=ctx.classe.cycle, nom='5e B', actif=True)
        ins2 = baker.make('inscriptions.Inscription', classe=autre_classe, annee_scolaire=ctx.annee, statut='ACTIF',
                          eleve=baker.make('inscriptions.Eleve', nom='SAWADOGO', prenom='Issa', photo=''))
        b2 = _bourse(ctx, ctx.fixe, '30000', inscription=ins2)
        baker.make('finances.BourseEleve', inscription=ins2, type_bourse=ctx.fixe, montant_accorde=Decimal('99999'),
                   actif=False, is_active=True)
        return b1, b2

    def test_liste_filtres_et_total(self, ctx, bourses):
        b1, b2 = bourses
        url = reverse('finances:boursiers_list')
        reponse = ctx.client.get(url)
        assert reponse.status_code == 200
        assert list(reponse.context['bourses']) == [b2, b1], 'tri par classe puis nom ; bourse désactivée ignorée'
        assert reponse.context['total_reduction'] == Decimal('55000')
        assert list(ctx.client.get(url, {'classe': ctx.classe.pk}).context['bourses']) == [b1]
        reponse = ctx.client.get(url, {'source': 'ONG'})
        assert list(reponse.context['bourses']) == [b2] and reponse.context['total_reduction'] == Decimal('30000')
        assert reponse.context['annee_sel'] == ctx.annee

    def test_pdf(self, ctx, bourses, mocker):
        weasy = mocker.patch('finances.views.bourses.WeasyHTML')
        weasy.return_value.write_pdf.side_effect = lambda buffer: buffer.write(b'%PDF-1.7')
        reponse = ctx.client.get(reverse('finances:boursiers_pdf'), {'source': 'ETABLISSEMENT'})
        assert reponse.status_code == 200 and reponse['Content-Type'] == 'application/pdf'
        html = weasy.call_args.kwargs['string']
        assert 'OUEDRAOGO' in html and 'SAWADOGO' not in html


@pytest.mark.django_db
class TestTypeBourse:
    def test_liste(self, ctx):
        reponse = ctx.client.get(reverse('finances:type_bourse_list'))
        assert reponse.status_code == 200 and set(reponse.context['types']) == {ctx.pct, ctx.fixe}

    def test_creation_valide(self, ctx):
        url = reverse('finances:type_bourse_create')
        assert ctx.client.get(url).status_code == 200
        reponse = ctx.client.post(url, {'nom': 'Mérite', 'code': 'mer', 'source': 'ETAT',
                                        'type_reduction': 'POURCENTAGE', 'valeur_reduction': '50',
                                        'rubrique': ctx.rubrique.pk, 'actif': 'on'})
        assert reponse.status_code == 302 and reponse.url == reverse('finances:type_bourse_list')
        tb = TypeBourse.objects.get(code='MER')
        assert (tb.nom, tb.source, tb.valeur_reduction, tb.rubrique, tb.actif) == ('Mérite', 'ETAT', Decimal('50'), ctx.rubrique, True)
        assert tb.created_by == ctx.user

    @pytest.mark.parametrize('donnees, erreur', [
        ({'nom': '', 'code': 'X1', 'valeur_reduction': '10'}, 'Le nom est obligatoire.'),
        ({'nom': 'X', 'code': '', 'valeur_reduction': '10'}, 'Le code est obligatoire.'),
        ({'nom': 'X', 'code': 'X1', 'valeur_reduction': '150'}, 'Un pourcentage ne peut pas depasser 100.'),
        ({'nom': 'X', 'code': 'X1', 'valeur_reduction': '0'}, 'La valeur de reduction doit etre positive.'),
        ({'nom': 'X', 'code': 'X1', 'valeur_reduction': 'abc'}, 'Valeur de reduction invalide.'),
        ({'nom': 'X', 'code': 'exc', 'valeur_reduction': '10'}, "Le code 'EXC' est deja utilise."),
    ])
    def test_creation_refusee(self, ctx, donnees, erreur):
        reponse = ctx.client.post(reverse('finances:type_bourse_create'), {'type_reduction': 'POURCENTAGE', **donnees})
        assert reponse.status_code == 200 and erreur in _messages(reponse)
        assert TypeBourse.objects.count() == 2

    def test_modification(self, ctx):
        url = reverse('finances:type_bourse_edit', kwargs={'type_id': ctx.fixe.pk})
        assert ctx.client.get(url).context['instance'] == ctx.fixe
        reponse = ctx.client.post(url, {'nom': 'Aide ONG', 'code': 'ONG1', 'source': 'ONG',
                                        'type_reduction': 'MONTANT_FIXE', 'valeur_reduction': '35000', 'actif': 'on'})
        assert reponse.status_code == 302
        ctx.fixe.refresh_from_db()
        assert ctx.fixe.valeur_reduction == Decimal('35000') and ctx.fixe.updated_by == ctx.user

    def test_desactivation(self, ctx):
        _bourse(ctx, ctx.pct, '10000')
        url = reverse('finances:type_bourse_delete', kwargs={'type_id': ctx.pct.pk})
        reponse = ctx.client.post(url, {'motif': 'Fin du programme'})
        ctx.pct.refresh_from_db()
        assert ctx.pct.actif and 'Impossible de désactiver : des bourses actives utilisent ce type.' in _messages(reponse)

        url = reverse('finances:type_bourse_delete', kwargs={'type_id': ctx.fixe.pk})
        reponse = ctx.client.post(url, {'motif': ''})
        ctx.fixe.refresh_from_db()
        assert ctx.fixe.actif and 'Le motif de désactivation est obligatoire.' in _messages(reponse)
        ctx.client.post(url, {'motif': 'Fin du programme'})
        ctx.fixe.refresh_from_db()
        assert (ctx.fixe.actif, ctx.fixe.is_active) == (False, False)


@pytest.mark.django_db
class TestEcheancierIndividuel:
    def test_creation(self, ctx):
        url = reverse('finances:echeancier_create', kwargs={'inscription_id': ctx.inscription.pk})
        assert ctx.client.get(url).context['title'] == 'Créer une échéance'
        reponse = ctx.client.post(url, {'libelle': '1re tranche', 'date_limite': '2026-11-15', 'montant_du': '50000'})
        assert reponse.status_code == 302
        ech = Echeancier.objects.get()
        assert (ech.inscription, ech.libelle, ech.date_limite, ech.montant_du, ech.paye) == \
            (ctx.inscription, '1re tranche', date(2026, 11, 15), Decimal('50000'), False)
        assert ech.created_by == ctx.user

    def test_creation_invalide(self, ctx):
        url = reverse('finances:echeancier_create', kwargs={'inscription_id': ctx.inscription.pk})
        reponse = ctx.client.post(url, {'libelle': '', 'date_limite': 'pas-une-date', 'montant_du': ''})
        assert reponse.status_code == 200
        assert set(reponse.context['form'].errors) == {'libelle', 'date_limite', 'montant_du'}
        assert not Echeancier.objects.exists()

    def test_modification_marque_payee(self, ctx):
        ech = baker.make('finances.Echeancier', inscription=ctx.inscription, libelle='2e tranche',
                         date_limite=date(2027, 1, 15), montant_du=Decimal('40000'), paye=False, is_active=True)
        url = reverse('finances:echeancier_edit', kwargs={'echeancier_id': ech.pk})
        assert ctx.client.get(url).context['echeancier'] == ech
        reponse = ctx.client.post(url, {'libelle': '2e tranche', 'date_limite': '2027-01-20',
                                        'montant_du': '40000', 'paye': 'on'})
        assert reponse.status_code == 302
        ech.refresh_from_db()
        assert (ech.paye, ech.date_limite, ech.updated_by) == (True, date(2027, 1, 20), ctx.user)

    def test_desactivation_exige_un_motif(self, ctx):
        ech = baker.make('finances.Echeancier', inscription=ctx.inscription, libelle='3e tranche',
                         date_limite=date(2027, 3, 1), montant_du=Decimal('30000'), is_active=True)
        url = reverse('finances:echeancier_delete', kwargs={'echeancier_id': ech.pk})
        assert ctx.client.get(url).context['echeancier'] == ech
        reponse = ctx.client.post(url, {'motif': ''})
        ech.refresh_from_db()
        assert ech.is_active and 'Le motif de désactivation est obligatoire.' in _messages(reponse)
        ctx.client.post(url, {'motif': 'Doublon'})
        ech.refresh_from_db()
        assert ech.is_active is False and Echeancier.objects.filter(pk=ech.pk).exists()
        assert ctx.client.get(url).status_code == 404, 'une échéance désactivée n’est plus modifiable'


@pytest.mark.django_db
class TestEcheancierGlobal:
    @pytest.fixture
    def echeances(self, ctx):
        today = timezone.now().date()

        def faire(libelle, jours, montant, paye=False, **extra):
            return baker.make('finances.Echeancier', inscription=ctx.inscription, libelle=libelle,
                              date_limite=today + timedelta(days=jours), montant_du=Decimal(montant), paye=paye,
                              is_active=True, **extra)

        retard, proche = faire('Retard', -1, '10000'), faire('Proche', 3, '20000')
        avenir, reglee = faire('Plus tard', 30, '30000'), faire('Réglée', -10, '40000', paye=True)
        baker.make('finances.Echeancier', inscription=ctx.inscription, libelle='Supprimée', date_limite=today,
                   montant_du=Decimal('99999'), is_active=False)
        return SimpleNamespace(retard=retard, proche=proche, avenir=avenir, reglee=reglee)

    def test_compteurs_totaux_et_ordre(self, ctx, echeances):
        e = echeances
        reponse = ctx.client.get(reverse('finances:echeancier_global'))
        assert reponse.status_code == 200
        c = reponse.context
        assert (c['nb_retard'], c['nb_alerte'], c['nb_avenir'], c['nb_paye']) == (1, 1, 1, 1)
        assert (c['total_du'], c['total_paye']) == (Decimal('60000'), Decimal('40000'))
        assert list(c['page_obj'].object_list) == [e.reglee, e.retard, e.proche, e.avenir], 'tri par date limite'

    def test_filtres(self, ctx, echeances):
        e = echeances
        url = reverse('finances:echeancier_global')
        assert list(ctx.client.get(url, {'statut': 'retard'}).context['page_obj'].object_list) == [e.retard]
        assert list(ctx.client.get(url, {'statut': 'alerte'}).context['page_obj'].object_list) == [e.proche]
        assert list(ctx.client.get(url, {'statut': 'avenir'}).context['page_obj'].object_list) == [e.avenir]
        assert list(ctx.client.get(url, {'statut': 'paye'}).context['page_obj'].object_list) == [e.reglee]
        assert list(ctx.client.get(url, {'q': 'Proche'}).context['page_obj'].object_list) == [e.proche]
        assert ctx.client.get(url, {'q': 'ouedraogo'}).context['page_obj'].paginator.count == 4
        autre_classe = baker.make('parametres.Classe', etablissement=ctx.etab, cycle=ctx.classe.cycle, actif=True)
        assert ctx.client.get(url, {'classe': autre_classe.pk}).context['page_obj'].paginator.count == 0

    def test_sans_etablissement_redirige(self, client):
        client.force_login(baker.make('accounts.User', is_superuser=True, role='SUPER_ADMIN', etablissement=None))
        reponse = client.get(reverse('finances:echeancier_global'))
        assert reponse.status_code == 302 and reponse.url == reverse('finances:paiement_list')

    def test_export_excel(self, ctx, echeances):
        reponse = ctx.client.get(reverse('finances:echeancier_global_xlsx'))
        assert reponse.status_code == 200
        assert reponse['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        assert 'echeanciers_2026-2027.xlsx' in reponse['Content-Disposition']
        feuille = load_workbook(io.BytesIO(reponse.content)).active
        lignes = {row[4]: row for row in feuille.iter_rows(values_only=True) if row and row[4] in
                  ('Retard', 'Proche', 'Plus tard', 'Réglée', 'Supprimée')}
        assert {lib: row[7] for lib, row in lignes.items()} == \
            {'Retard': 'En retard', 'Proche': 'Échéance proche', 'Plus tard': 'À venir', 'Réglée': 'Payé'}
        assert lignes['Retard'][:3] == ('OUEDRAOGO Awa', 'M-001', '6e A') and lignes['Retard'][6] == 10000.0

        reponse = ctx.client.get(reverse('finances:echeancier_global_xlsx'), {'statut': 'paye'})
        feuille = load_workbook(io.BytesIO(reponse.content)).active
        assert [row[4] for row in feuille.iter_rows(values_only=True) if row and row[7] in ('Payé', 'En retard')] == ['Réglée']
